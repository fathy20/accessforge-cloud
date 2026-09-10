"""The member-duty Heavy allowance model, pinned to the July 2026 evidence.

Every case here replicates a real member shape from the manual reference
workbook that validated the model 54/55 (names replaced by codes).
"""

import unittest

from backend.statistics.crew_hours.allowance import (
    CREDIT_LEON,
    CREDIT_SWAP,
    AllowanceLeg,
    compute_member_credits,
)


def leg(key, date, start, end, position, *, leon=None, adep="HRG", ades="SSH"):
    return AllowanceLeg(
        key=key,
        flight_date=date,
        start_time=start,
        end_time=end,
        position=position,
        leon_heavy=leon,
        departure_airport=adep,
        arrival_airport=ades,
    )


class TestSwapCredit(unittest.TestCase):
    def test_operate_out_ride_pad_back_is_one_credit_painted_on_both_legs(self):
        # Donia's shape: FO out, PAD home, one duty -> exactly 1, both legs lit.
        result = compute_member_credits([
            leg("a", "10-06-2026", "14:25", "20:40", "FO", adep="HRG", ades="LIS"),
            leg("b", "10-06-2026", "21:50", "03:35", "PAD", adep="LIS", ades="HRG"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].source, CREDIT_SWAP)
        self.assertEqual(result.by_leg["a"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["b"], (True, CREDIT_SWAP))

    def test_ride_first_then_operate_is_the_same_credit(self):
        # Ziad Mohab's shape: PAD out, CPT home (direction is irrelevant).
        result = compute_member_credits([
            leg("a", "02-07-2026", "14:25", "19:50", "PAD", adep="HRG", ades="VKO"),
            leg("b", "02-07-2026", "21:20", "03:20", "CPT", adep="VKO", ades="HRG"),
        ])
        self.assertEqual(result.credits, 1)

    def test_two_swap_duties_are_two_credits(self):
        # John Ashraf's July: exactly 2 in the sheet, from 2 swap duties.
        result = compute_member_credits([
            leg("a1", "15-07-2026", "12:10", "12:55", "FO"),
            leg("a2", "15-07-2026", "14:10", "20:35", "FO", adep="HRG", ades="LIS"),
            leg("a3", "15-07-2026", "22:30", "04:10", "PAD", adep="LIS", ades="HRG"),
            leg("b1", "28-07-2026", "19:05", "00:55", "FO", adep="SSH", ades="ALA"),
            leg("b2", "29-07-2026", "02:10", "08:45", "PAD", adep="ALA", ades="SSH"),
        ])
        self.assertEqual(result.credits, 2)

    def test_operate_both_legs_without_riding_is_no_credit(self):
        # The 23-06 pair the owner asked about (RSX8891+RSX6083): the sheet
        # gave the all-operating member nothing for the identical July shape.
        # An operated leg CAN partner another operated leg (owner ruling
        # 2026-09-02) -- but this one's only neighbour, HRG->SSH, is domestic,
        # and a domestic leg is never a valid partner. No genuine return
        # exists in this duty, so it stays uncredited.
        result = compute_member_credits([
            leg("a", "23-06-2026", "05:30", "06:10", "FO", adep="HRG", ades="SSH"),
            leg("b", "23-06-2026", "07:00", "13:00", "FO", adep="SSH", ades="OPO"),
        ])
        self.assertEqual(result.credits, 0)
        self.assertIn("no other leg", result.duties[0].reason)

    def test_pad_only_duty_is_no_credit(self):
        result = compute_member_credits([
            leg("a", "01-07-2026", "00:00", "05:40", "PAD", adep="SVX", ades="SSH"),
        ])
        self.assertEqual(result.credits, 0)
        self.assertIn("rode PAD only", result.duties[0].reason)

    def test_an_operated_round_trip_with_no_ride_is_one_credit_on_both_legs(self):
        # The owner's governing example, verbatim: Cairo to a place in Russia
        # and back. Both legs are OPERATED -- nobody rides -- and both are
        # Heavy anyway, because the credit is about a genuine return, not
        # about who happened to be a passenger on it (owner ruling
        # 2026-09-02).
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "11:30", "FO", adep="CAI", ades="VKO"),
            leg("b", "12-07-2026", "13:00", "18:30", "FO", adep="VKO", ades="CAI"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].source, CREDIT_SWAP)
        self.assertEqual(result.by_leg["a"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["b"], (True, CREDIT_SWAP))

    def test_23_06_becomes_heavy_once_a_genuine_return_leg_exists(self):
        # The exact 23-06 duty (HRG->SSH shuttle, SSH->OPO 6:00), continued:
        # the owner's own follow-up to that case -- "you'll see the one after
        # it, he returns too" -- adds an OPO->SSH return within the link
        # break. The shuttle stays No (domestic); the two international legs
        # of the genuine out-and-back are both Heavy.
        result = compute_member_credits([
            leg("shuttle", "23-06-2026", "05:30", "06:10", "FO", adep="HRG", ades="SSH"),
            leg("out", "23-06-2026", "07:00", "13:00", "FO", adep="SSH", ades="OPO"),
            leg("back", "23-06-2026", "14:15", "20:00", "FO", adep="OPO", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.by_leg["shuttle"], (False, None))
        self.assertEqual(result.by_leg["out"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["back"], (True, CREDIT_SWAP))

    def test_two_short_operated_legs_neither_over_the_minimum_is_no_credit(self):
        # Karim Fekry's shape, generalized: even with a genuine (non-domestic)
        # partner nearby, neither operated leg exceeds the sector minimum, so
        # there is no qualifying anchor to credit from either direction.
        result = compute_member_credits([
            leg("a", "11-07-2026", "09:40", "12:50", "FO", adep="SSH", ades="KRR"),
            leg("b", "11-07-2026", "15:00", "18:20", "FO", adep="KRR", ades="SSH"),
        ])
        self.assertEqual(result.credits, 0)


class TestLeonCredit(unittest.TestCase):
    def test_leon_augmented_pair_is_one_credit_per_duty_not_per_leg(self):
        # The CGN 3-pilot rotations: LEON marks both legs; the sheet pays 1.
        result = compute_member_credits([
            leg("a", "10-07-2026", "17:05", "22:15", "CPT2", leon=True, adep="HRG", ades="CGN"),
            leg("b", "10-07-2026", "23:40", "04:40", "CPT2", leon=True, adep="CGN", ades="HRG"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].source, CREDIT_LEON)

    def test_leon_false_never_credits_by_itself(self):
        result = compute_member_credits([
            leg("a", "13-07-2026", "14:30", "20:50", "FO", leon=False, adep="HRG", ades="OPO"),
        ])
        self.assertEqual(result.credits, 0)


class TestNeutralSlots(unittest.TestCase):
    def test_a_short_psn_shuttle_neither_operates_nor_rides(self):
        # Karim Fekry's 11-07: a 0:40 PSN base shuttle + operated KRR pair ->
        # NOT credited (the July sheet paid him nothing for this duty).
        result = compute_member_credits([
            leg("a", "11-07-2026", "06:35", "07:15", "PSN"),
            leg("b", "11-07-2026", "09:40", "12:50", "FO", adep="SSH", ades="KRR"),
            leg("c", "11-07-2026", "15:00", "18:20", "FO", adep="KRR", ades="SSH"),
        ])
        self.assertEqual(result.credits, 0)

    def test_psn_on_a_rotation_scale_sector_rides(self):
        # Owner case 29-06 (Ahmed Kamel): PSN HRG->OPO 5:45 out, FO back ->
        # the swap credit, painted on both legs.
        result = compute_member_credits([
            leg("a", "29-06-2026", "14:35", "20:20", "PSN", adep="HRG", ades="OPO"),
            leg("b", "29-06-2026", "21:25", "03:10", "FO", adep="OPO", ades="HRG"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].source, CREDIT_SWAP)
        self.assertEqual(result.by_leg["a"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["b"], (True, CREDIT_SWAP))

    def test_operate_then_long_psn_back_rides_too(self):
        # Owner case 09-06 (Ahmed Tahoon): CPT out HRG->OPO, PSN 5:20 back.
        result = compute_member_credits([
            leg("a", "09-06-2026", "07:15", "13:20", "CPT", adep="HRG", ades="OPO"),
            leg("b", "09-06-2026", "14:45", "20:05", "PSN", adep="OPO", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)

    def test_obs_pair_is_no_credit_even_when_a_rotation_exists(self):
        # Amro Nasef's OGZ pair: sheet blank; observers never earn.
        result = compute_member_credits([
            leg("a", "13-07-2026", "09:25", "12:35", "OBS", adep="SSH", ades="OGZ"),
            leg("b", "13-07-2026", "14:50", "18:20", "OBS", adep="OGZ", ades="SSH"),
        ])
        self.assertEqual(result.credits, 0)

    def test_trainee_slots_earn_nothing_without_any_rank_rule(self):
        result = compute_member_credits([
            leg("a", "05-07-2026", "10:00", "12:00", "SP"),
            leg("b", "05-07-2026", "13:00", "15:00", "OPS"),
        ])
        self.assertEqual(result.credits, 0)


class TestDomesticVeto(unittest.TestCase):
    """Owner ruling 2026-09-02: inside Egypt is never Heavy, and its hours
    never add up to reach the sector minimum."""

    def test_a_long_domestic_sector_with_a_ride_is_not_heavy(self):
        # CAI->SSH is domestic: over 4h and paired with a ride, still nothing.
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "10:30", "FO", adep="CAI", ades="SSH"),
            leg("b", "12-07-2026", "12:00", "16:30", "PAD", adep="SSH", ades="CAI"),
        ])
        self.assertEqual(result.credits, 0)

    def test_domestic_hours_never_accumulate_to_reach_the_minimum(self):
        # Three domestic sectors totalling 9:00 in one duty — still nothing.
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "09:00", "FO", adep="CAI", ades="SSH"),
            leg("b", "12-07-2026", "10:00", "13:00", "FO", adep="SSH", ades="HRG"),
            leg("c", "12-07-2026", "14:00", "17:00", "PAD", adep="HRG", ades="CAI"),
        ])
        self.assertEqual(result.credits, 0)

    def test_the_icao_spelling_of_an_egyptian_airport_is_domestic_too(self):
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "10:30", "FO", adep="HECA", ades="HESH"),
            leg("b", "12-07-2026", "12:00", "16:30", "PAD", adep="HESH", ades="HECA"),
        ])
        self.assertEqual(result.credits, 0)

    def test_cairo_to_russia_and_back_is_heavy(self):
        # The owner's governing example: one end outside Egypt makes it a
        # rotation, and it clears the sector minimum.
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "11:30", "FO", adep="CAI", ades="VKO"),
            leg("b", "12-07-2026", "13:00", "18:30", "PAD", adep="VKO", ades="CAI"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].source, CREDIT_SWAP)

    def test_a_domestic_hop_inside_a_credited_duty_is_painted_no(self):
        # The owner's 23-06 shape, completed with a ride: the shuttle reads No,
        # the international legs of the same credited duty read Yes.
        result = compute_member_credits([
            leg("shuttle", "23-06-2026", "05:30", "06:10", "FO", adep="HRG", ades="SSH"),
            leg("out", "23-06-2026", "07:00", "13:00", "FO", adep="SSH", ades="OPO"),
            leg("back", "23-06-2026", "14:15", "20:00", "PAD", adep="OPO", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.by_leg["shuttle"], (False, None))
        self.assertEqual(result.by_leg["out"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["back"], (True, CREDIT_SWAP))

    def test_a_leon_credited_domestic_leg_keeps_painting_yes(self):
        # LEON's own value is authoritative — the domestic carve-out only
        # applies to the local swap rule, never to CREDIT_LEON.
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "06:40", "CPT2", leon=True,
                adep="CAI", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.by_leg["a"], (True, CREDIT_LEON))

    def test_a_missing_airport_never_demotes_a_sector_to_domestic(self):
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "11:30", "FO", adep="CAI", ades=None),
            leg("b", "12-07-2026", "13:00", "18:30", "PAD", adep=None, ades="CAI"),
        ])
        self.assertEqual(result.credits, 1)


class TestSectorMinimum(unittest.TestCase):
    def test_an_international_sector_at_exactly_four_hours_does_not_qualify(self):
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "10:00", "FO", adep="CAI", ades="VKO"),
            leg("b", "12-07-2026", "11:00", "15:00", "PAD", adep="VKO", ades="CAI"),
        ])
        self.assertEqual(result.credits, 0)

    def test_one_minute_over_four_hours_qualifies(self):
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "10:01", "FO", adep="CAI", ades="VKO"),
            leg("b", "12-07-2026", "11:00", "15:00", "PAD", adep="VKO", ades="CAI"),
        ])
        self.assertEqual(result.credits, 1)

    def test_leon_augmented_ignores_the_sector_minimum_and_the_domestic_veto(self):
        # LEON's own value is authoritative and is never re-judged locally
        # (owner ruling 2026-09-02: "LEON pulls it correctly").
        result = compute_member_credits([
            leg("a", "12-07-2026", "06:00", "06:40", "CPT2", leon=True,
                adep="CAI", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].source, CREDIT_LEON)


class TestMonthBoundary(unittest.TestCase):
    """A rotation that straddles a month end is heavy on BOTH sheets, but is
    counted once — against the month of its first sector."""

    LEGS = [
        leg("out", "31-07-2026", "20:00", "01:30", "FO", adep="CAI", ades="VKO"),
        leg("back", "01-08-2026", "03:00", "08:30", "PAD", adep="VKO", ades="CAI"),
    ]

    def test_both_legs_read_heavy_when_the_window_is_the_departing_month(self):
        result = compute_member_credits(
            self.LEGS, window_start="2026-07-01", window_end="2026-07-31"
        )
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.by_leg["out"][0], True)
        self.assertEqual(result.by_leg["back"][0], True)

    def test_both_legs_still_read_heavy_on_the_following_month_sheet(self):
        result = compute_member_credits(
            self.LEGS, window_start="2026-08-01", window_end="2026-08-31"
        )
        self.assertEqual(result.by_leg["out"][0], True)
        self.assertEqual(result.by_leg["back"][0], True)

    def test_the_credit_is_counted_once_in_the_departing_month_only(self):
        july = compute_member_credits(
            self.LEGS, window_start="2026-07-01", window_end="2026-07-31"
        )
        august = compute_member_credits(
            self.LEGS, window_start="2026-08-01", window_end="2026-08-31"
        )
        self.assertEqual((july.credits, august.credits), (1, 0))


class TestEvnVeto(unittest.TestCase):
    def test_evn_legs_contribute_nothing_in_either_role(self):
        # Operated + PAD across EVN sectors: still nothing (owner absolute).
        result = compute_member_credits([
            leg("a", "01-07-2026", "23:05", "02:00", "FO", adep="SSH", ades="EVN"),
            leg("b", "02-07-2026", "03:10", "06:30", "PAD", adep="EVN", ades="SSH"),
        ])
        self.assertEqual(result.credits, 0)

    def test_evn_in_icao_form_is_the_same_airport(self):
        result = compute_member_credits([
            leg("a", "01-07-2026", "23:05", "02:00", "FO", adep="HESH", ades="UDYZ"),
            leg("b", "02-07-2026", "03:10", "06:30", "PAD", adep="UDYZ", ades="HESH"),
        ])
        self.assertEqual(result.credits, 0)


class TestEvnPainting(unittest.TestCase):
    """Owner ruling 2026-09-10: an EVN leg reads No on every surface, even
    inside a duty that another leg credited. The duty keeps its credit and the
    H.C count never moves — this is display painting only."""

    def test_evn_inside_a_leon_credited_duty_is_painted_no_and_the_count_holds(self):
        # The real MSN 06/07-06 shape: a LEON-augmented CAI->SSH sector, then
        # the SSH<->EVN pair inside the same duty. EVN beats LEON here, unlike
        # the domestic carve-out which is swap-only.
        result = compute_member_credits([
            leg("earner", "06-06-2026", "17:55", "19:05", "FA2", leon=True,
                adep="CAI", ades="SSH"),
            leg("evn_out", "06-06-2026", "21:45", "00:25", "IFA",
                adep="SSH", ades="EVN"),
            leg("evn_back", "07-06-2026", "01:40", "04:40", "IFA",
                adep="EVN", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.by_leg["earner"], (True, CREDIT_LEON))
        self.assertEqual(result.by_leg["evn_out"], (False, None))
        self.assertEqual(result.by_leg["evn_back"], (False, None))

    def test_leon_marking_the_evn_legs_augmented_does_not_rescue_them(self):
        # The literal all_legs.csv shape: LEON reports leon_heavy=True on the
        # EVN legs themselves (06-06 RSX121 SSH-EVN, 07-06 RSX122 EVN-SSH; the
        # 04/05-07 AMY pair is identical). The EVN veto runs before LEON's
        # value is ever read, so the earning path is still the non-EVN CAI->SSH
        # sector -- and the EVN rows read No even though LEON called them
        # augmented. This is the strongest form of the carve-out.
        result = compute_member_credits([
            leg("earner", "06-06-2026", "17:55", "19:05", "FA2", leon=True,
                adep="CAI", ades="SSH"),
            leg("evn_out", "06-06-2026", "21:45", "00:25", "IFA", leon=True,
                adep="SSH", ades="EVN"),
            leg("evn_back", "07-06-2026", "01:40", "04:40", "IFA", leon=True,
                adep="EVN", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)
        duty = result.duties[0]
        self.assertTrue(duty.credited)
        self.assertEqual(duty.source, CREDIT_LEON)
        self.assertEqual(result.by_leg["earner"], (True, CREDIT_LEON))
        self.assertEqual(result.by_leg["evn_out"], (False, None))
        self.assertEqual(result.by_leg["evn_back"], (False, None))

    def test_an_evn_only_duty_leon_marked_augmented_still_earns_nothing(self):
        # Same LEON value, no non-EVN sector to earn from: the veto removes the
        # only legs there are, so there is no credit to paint at all.
        result = compute_member_credits([
            leg("evn_out", "06-06-2026", "21:45", "00:25", "IFA", leon=True,
                adep="SSH", ades="EVN"),
            leg("evn_back", "07-06-2026", "01:40", "04:40", "IFA", leon=True,
                adep="EVN", ades="SSH"),
        ])
        self.assertEqual(result.credits, 0)
        self.assertEqual(result.by_leg["evn_out"], (False, None))
        self.assertEqual(result.by_leg["evn_back"], (False, None))

    def test_evn_inside_a_swap_credited_duty_is_painted_no_and_the_count_holds(self):
        result = compute_member_credits([
            leg("out", "12-07-2026", "06:00", "11:30", "FO", adep="CAI", ades="VKO"),
            leg("back", "12-07-2026", "13:00", "18:30", "PAD", adep="VKO", ades="CAI"),
            leg("evn", "12-07-2026", "20:00", "22:40", "FO", adep="SSH", ades="EVN"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].source, CREDIT_SWAP)
        self.assertEqual(result.by_leg["out"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["back"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["evn"], (False, None))

    def test_the_icao_spelling_of_evn_is_carved_out_identically(self):
        result = compute_member_credits([
            leg("earner", "06-06-2026", "17:55", "19:05", "FA2", leon=True,
                adep="CAI", ades="SSH"),
            leg("evn", "06-06-2026", "21:45", "00:25", "IFA",
                adep="HESH", ades="UDYZ"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.by_leg["earner"], (True, CREDIT_LEON))
        self.assertEqual(result.by_leg["evn"], (False, None))

    def test_an_evn_pad_leg_and_an_evn_psn_leg_are_both_painted_no(self):
        result = compute_member_credits([
            leg("earner", "06-06-2026", "17:55", "19:05", "FA2", leon=True,
                adep="CAI", ades="SSH"),
            leg("evn_pad", "06-06-2026", "21:45", "00:25", "PAD",
                adep="SSH", ades="EVN"),
            leg("evn_psn", "07-06-2026", "01:40", "04:40", "PSN",
                adep="EVN", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.by_leg["evn_pad"], (False, None))
        self.assertEqual(result.by_leg["evn_psn"], (False, None))

    def test_the_psn_chaining_of_neighbouring_legs_survives_an_evn_neighbour(self):
        # The PSN carve-out reads its neighbours from the duty list, never from
        # by_leg, so carving the EVN leg out cannot break the chain that the
        # non-EVN PSN leg depends on: 20:05 -> 21:10 is 1:05, still chained.
        result = compute_member_credits([
            leg("out", "09-06-2026", "07:15", "13:20", "CPT", adep="HRG", ades="OPO"),
            leg("psn_back", "09-06-2026", "14:45", "20:05", "PSN",
                adep="OPO", ades="SSH"),
            leg("evn", "09-06-2026", "21:10", "23:55", "FO", adep="SSH", ades="EVN"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].source, CREDIT_SWAP)
        self.assertEqual(result.by_leg["out"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["psn_back"], (True, CREDIT_SWAP))
        self.assertEqual(result.by_leg["evn"], (False, None))

    def test_the_duty_credit_record_still_reports_its_source_and_reason(self):
        result = compute_member_credits([
            leg("earner", "06-06-2026", "17:55", "19:05", "FA2", leon=True,
                adep="CAI", ades="SSH"),
            leg("evn", "06-06-2026", "21:45", "00:25", "IFA",
                adep="SSH", ades="EVN"),
        ])
        duty = result.duties[0]
        self.assertTrue(duty.credited)
        self.assertEqual(duty.source, CREDIT_LEON)
        self.assertEqual(duty.reason, "LEON marked an operated sector augmented")
        self.assertEqual(duty.leg_keys, ("earner", "evn"))

    def test_an_uncredited_duty_paints_its_evn_leg_no_the_same_way(self):
        # The carve-out is guarded on `earned`, so an uncredited duty is
        # untouched: its EVN leg reads No because the duty does.
        result = compute_member_credits([
            leg("evn", "06-06-2026", "21:45", "00:25", "IFA",
                adep="SSH", ades="EVN"),
        ])
        self.assertEqual(result.credits, 0)
        self.assertEqual(result.by_leg["evn"], (False, None))


class TestDutyBoundaries(unittest.TestCase):
    def test_a_four_hour_break_splits_the_duty_strictly(self):
        # 4:00 exactly rejects -> two duties, each missing the other role.
        result = compute_member_credits([
            leg("a", "10-07-2026", "08:00", "10:00", "FO"),
            leg("b", "10-07-2026", "14:00", "16:00", "PAD"),
        ])
        self.assertEqual(result.credits, 0)
        self.assertEqual(len(result.duties), 2)

    def test_a_359_break_keeps_one_duty(self):
        # 3:59 is inside the 4h duty limit, so this is ONE duty — but the ride
        # sits further than SWAP_LINK_BREAK from the operated sector, so the
        # duty earns nothing. One duty, no credit.
        result = compute_member_credits([
            leg("a", "10-07-2026", "08:00", "14:30", "FO", adep="HRG", ades="LIS"),
            leg("b", "10-07-2026", "18:29", "23:00", "PAD", adep="LIS", ades="HRG"),
        ])
        self.assertEqual(len(result.duties), 1)
        self.assertEqual(result.credits, 0)

    def test_a_ride_within_three_hours_of_the_operated_sector_credits(self):
        result = compute_member_credits([
            leg("a", "10-07-2026", "08:00", "14:30", "FO", adep="HRG", ades="LIS"),
            leg("b", "10-07-2026", "17:29", "22:00", "PAD", adep="LIS", ades="HRG"),
        ])
        self.assertEqual(len(result.duties), 1)
        self.assertEqual(result.credits, 1)

    def test_crossing_midnight_is_one_duty_not_two_days(self):
        result = compute_member_credits([
            leg("a", "16-06-2026", "17:15", "22:35", "FO", adep="SSH", ades="SVX"),
            leg("b", "16-06-2026", "23:50", "06:00", "PAD", adep="SVX", ades="SSH"),
        ])
        self.assertEqual(result.credits, 1)
        self.assertEqual(result.duties[0].anchor_utc_date, "2026-06-16")

    def test_duty_is_attributed_to_its_anchor_month(self):
        # Salah Hesham's 30-06 -> 01-07 SVX duty: July's sheet did not pay it.
        result = compute_member_credits(
            [
                leg("a", "30-06-2026", "17:15", "23:00", "FO", adep="SSH", ades="SVX"),
                leg("b", "01-07-2026", "00:00", "05:40", "PAD", adep="SVX", ades="SSH"),
            ],
            window_start="2026-07-01",
            window_end="2026-07-31",
        )
        self.assertEqual(result.credits, 0)
        self.assertIn("outside this window", result.duties[0].reason)
        # ...but the rotation IS heavy, so both legs read Yes on their sheets —
        # the July sheet simply does not pay the credit (owner ruling 2026-09-02).
        self.assertEqual(result.by_leg["a"][0], True)
        self.assertEqual(result.by_leg["b"][0], True)

    def test_a_duty_anchored_on_the_last_day_of_the_window_counts(self):
        result = compute_member_credits(
            [
                leg("a", "31-07-2026", "17:20", "22:45", "FO", adep="HRG", ades="LED"),
                leg("b", "31-07-2026", "23:50", "05:20", "PAD", adep="LED", ades="HRG"),
            ],
            window_start="2026-07-01",
            window_end="2026-07-31",
        )
        self.assertEqual(result.credits, 1)

    def test_iso_timestamps_are_accepted_too(self):
        result = compute_member_credits([
            leg("a", None, "2026-06-10T14:25:00Z", "2026-06-10T20:40:00Z", "FO",
                adep="HRG", ades="LIS"),
            leg("b", None, "2026-06-10T21:50:00Z", "2026-06-11T03:35:00Z", "PAD",
                adep="LIS", ades="HRG"),
        ])
        self.assertEqual(result.credits, 1)

    def test_unusable_times_exclude_the_leg_without_crashing(self):
        result = compute_member_credits([
            leg("a", None, None, None, "FO"),
            leg("b", "10-07-2026", "08:00", "10:00", "PAD"),
        ])
        self.assertEqual(result.credits, 0)
        self.assertNotIn("a", result.by_leg)


class TestServiceWiring(unittest.TestCase):
    """The report response carries H.C and paints both legs of a credited duty."""

    def _response(
        self,
        rows,
        buffered_rows=None,
        from_date="2026-06-01",
        to_date="2026-06-30",
    ):
        from backend.statistics.crew_hours.augmented import AugmentedIndex
        from backend.statistics.crew_hours.crew_context import CrewContextIndex
        from backend.statistics.crew_hours.mcp_report import OfficialMcpReport
        from backend.statistics.crew_hours.service import _build_mcp_report_response

        totals = {code: "10:00" for row in rows for code in row["crew_codes"]}
        return _build_mcp_report_response(
            OfficialMcpReport(totals, rows, buffered_rows=buffered_rows),
            from_date=from_date,
            to_date=to_date,
            position="All",
            crew_member=None,
            augmented_index=AugmentedIndex(True, {}, 0, 0, {}),
            crew_context_index=CrewContextIndex(False, {}),
        )

    @staticmethod
    def _row(uid, number, adep, ades, date, start, end, codes, positions):
        return {
            "scope_row_unique_id": f"row-{uid}",
            "unique_id": uid,
            "flightNo": number,
            "crew_codes": codes,
            "crew_names": [f"Crew {c}" for c in codes],
            "crew_position_names": positions,
            "acftType": "B738 - 737-800",
            "blockTimeJourneyLog": "01:30",
            "jl_adep_preferred_code": adep,
            "jl_ades_preferred_code": ades,
            "date_STD_log_UTC": date,
            "JL_STD_UTC": start,
            "JL_STA_UTC": end,
        }

    def test_the_owner_complaint_donia_both_legs_credited_and_hc_is_one(self):
        # 10-06: FO out HRG->LIS, PAD home LIS->HRG. The old per-leg verdict
        # split them (No / Yes+badge); the allowance paints BOTH and H.C = 1.
        rows = [
            self._row(901, "RSX6077", "HRG", "LIS", "10-06-2026", "14:25", "20:40",
                      ["DON"], ["FO"]),
            self._row(902, "RSX6078", "LIS", "HRG", "10-06-2026", "21:50", "03:35",
                      ["DON"], ["PAD"]),
        ]
        member = self._response(rows).crew_members[0]

        self.assertEqual(member.heavy_credits, 1)
        by_number = {f.flight_number: f for f in member.flights}
        for number in ("RSX6077", "RSX6078"):
            self.assertIs(by_number[number].duty_credit, True, number)
            self.assertEqual(by_number[number].credit_source, "OPERATE_PLUS_RIDE")

    def test_a_month_straddling_rotation_paints_the_new_month_leg(self):
        # The 31-07 outbound is narrowed out of an August pull's DISPLAYED rows
        # but survives in buffered_rows, so the duty chain still sees the whole
        # rotation: the 01-08 return paints Yes while the credit stays July's.
        outbound = self._row(920, "RSX6077", "HRG", "LIS", "31-07-2026",
                             "20:00", "02:15", ["DON"], ["FO"])
        ride_home = self._row(921, "RSX6078", "LIS", "HRG", "01-08-2026",
                              "03:30", "09:15", ["DON"], ["PAD"])

        august = self._response(
            [ride_home],
            buffered_rows=[outbound, ride_home],
            from_date="2026-08-01",
            to_date="2026-08-31",
        ).crew_members[0]
        self.assertEqual(august.flight_count, 1)  # display shows August only
        self.assertEqual(august.heavy_credits, 0)  # the credit belongs to July
        self.assertIs(august.flights[0].duty_credit, True)
        self.assertEqual(august.flights[0].credit_source, "OPERATE_PLUS_RIDE")

        july = self._response(
            [outbound],
            buffered_rows=[outbound, ride_home],
            from_date="2026-07-01",
            to_date="2026-07-31",
        ).crew_members[0]
        self.assertEqual(july.heavy_credits, 1)
        self.assertIs(july.flights[0].duty_credit, True)

    def test_the_owner_complaint_all_operating_pair_is_uncredited_for_both(self):
        # 23-06 RSX8891 + RSX6083: operated both, no PAD, LEON silent ->
        # symmetric No-credit (the July sheet paid nothing for this shape).
        rows = [
            self._row(911, "RSX8891", "HRG", "SSH", "23-06-2026", "05:30", "06:10",
                      ["C1"], ["FO"]),
            self._row(912, "RSX6083", "SSH", "OPO", "23-06-2026", "07:00", "13:00",
                      ["C1"], ["FO"]),
        ]
        member = self._response(rows).crew_members[0]

        self.assertEqual(member.heavy_credits, 0)
        for flight in member.flights:
            self.assertIs(flight.duty_credit, False, flight.flight_number)
            self.assertIsNone(flight.credit_source)


class TestEvnPaintingReachesTheScreenAndTheExport(unittest.TestCase):
    """The real MSN 06/07-06 case, end to end through LiveCrewHoursService and
    the XLSX writer: the two EVN rows display No, everything that is counted or
    totalled is byte-for-byte what it was before the carve-out existed."""

    ROWS = [
        {
            "scope_row_unique_id": "row-1101",
            "unique_id": 1101,
            "flightNo": "RSX492",
            "crew_codes": ["MSN"],
            "crew_names": ["Crew MSN"],
            "crew_position_names": ["FA2"],
            "acftType": "B738 - 737-800",
            "blockTimeJourneyLog": "01:30",
            "jl_adep_preferred_code": "CAI",
            "jl_ades_preferred_code": "SSH",
            "date_STD_log_UTC": "06-06-2026",
            "JL_STD_UTC": "17:55",
            "JL_STA_UTC": "19:05",
        },
        {
            "scope_row_unique_id": "row-1102",
            "unique_id": 1102,
            "flightNo": "RSX121",
            "crew_codes": ["MSN"],
            "crew_names": ["Crew MSN"],
            "crew_position_names": ["IFA"],
            "acftType": "B738 - 737-800",
            "blockTimeJourneyLog": "01:30",
            "jl_adep_preferred_code": "SSH",
            "jl_ades_preferred_code": "EVN",
            "date_STD_log_UTC": "06-06-2026",
            "JL_STD_UTC": "21:45",
            "JL_STA_UTC": "00:25",
        },
        {
            "scope_row_unique_id": "row-1103",
            "unique_id": 1103,
            "flightNo": "RSX122",
            "crew_codes": ["MSN"],
            "crew_names": ["Crew MSN"],
            "crew_position_names": ["IFA"],
            "acftType": "B738 - 737-800",
            "blockTimeJourneyLog": "01:30",
            "jl_adep_preferred_code": "EVN",
            "jl_ades_preferred_code": "SSH",
            "date_STD_log_UTC": "07-06-2026",
            "JL_STD_UTC": "01:40",
            "JL_STA_UTC": "04:40",
        },
    ]

    def _member(self):
        from backend.statistics.crew_hours.augmented import AugmentedIndex
        from backend.statistics.crew_hours.mcp_report import OfficialMcpReport
        from backend.statistics.crew_hours.service import LiveCrewHoursService

        rows = self.ROWS

        class FakeCrewClient:
            def fetch_official_totals(self, from_date, to_date):
                return OfficialMcpReport(
                    {"MSN": "10:00"}, rows, buffered_rows=rows
                )

            def fetch_augmented_index(self, from_date, to_date):
                # LEON marks the EVN pair augmented TOO, not just the CAI->SSH
                # earner: in the live pull the SSH<->EVN legs of this very
                # rotation (06-06 RSX121 SSH-EVN and 07-06 RSX122 EVN-SSH, and
                # the 04/05-07 SSH<->EVN pair) all carry leon_heavy=True --
                # 206 of the 474 EVN legs in the June+July pull do. That is the
                # real conflict the carve-out has to survive: LEON itself
                # called these EVN legs augmented, the credit still comes from
                # the non-EVN sector, and the owner absolute still paints the
                # EVN rows No.
                return AugmentedIndex(
                    True,
                    {("MSN", 1101): True, ("MSN", 1102): True, ("MSN", 1103): True},
                    3,
                    0,
                )

        report = LiveCrewHoursService(FakeCrewClient()).get_crew_hours_report(
            "2026-06-01", "2026-06-30"
        )
        return report, report.crew_members[0]

    def test_the_evn_rows_display_no_while_the_credit_and_the_count_stay_put(self):
        _, member = self._member()

        # Unchanged from before the carve-out: the duty still earns, and H.C
        # still reads 1 for this member.
        self.assertEqual(member.heavy_credits, 1)
        self.assertEqual(member.official_total, "10:00")
        self.assertEqual(member.flight_count, 3)

        by_number = {flight.flight_number: flight for flight in member.flights}
        # The earning leg is untouched.
        self.assertIs(by_number["RSX492"].duty_credit, True)
        self.assertEqual(by_number["RSX492"].credit_source, CREDIT_LEON)
        # Both EVN rows now read No, and claim no source.
        for number in ("RSX121", "RSX122"):
            self.assertIs(by_number[number].duty_credit, False, number)
            self.assertIsNone(by_number[number].credit_source, number)
        # Nothing else about the EVN rows moved.
        for number in ("RSX492", "RSX121", "RSX122"):
            self.assertEqual(by_number[number].block_time, "01:30", number)

    def test_the_evn_rows_carry_leons_yes_and_still_resolve_to_a_conflicted_no(self):
        # Proof that this fixture reproduces the live conflict rather than a
        # LEON-silent stand-in: the EVN rows arrive with leon_heavy=True and
        # still come out effective_heavy=False, flagged as a conflict, on the
        # EVN_AIRPORT absolute. The earner keeps LEON's Yes untouched.
        _, member = self._member()
        by_number = {flight.flight_number: flight for flight in member.flights}

        earner = by_number["RSX492"]
        self.assertIs(earner.leon_heavy, True)
        self.assertIs(earner.effective_heavy, True)
        self.assertFalse(earner.heavy_conflict)

        for number in ("RSX121", "RSX122"):
            row = by_number[number]
            self.assertIs(row.leon_heavy, True, number)
            self.assertIs(row.effective_heavy, False, number)
            self.assertTrue(row.heavy_conflict, number)
            self.assertEqual(row.heavy_source, "LOCAL_RULE", number)
            self.assertEqual(row.heavy_reason, "EVN_AIRPORT", number)

    def test_the_xlsx_shows_no_on_the_evn_rows_and_leaves_hc_and_totals_alone(self):
        import io
        from datetime import datetime, time, timedelta, timezone

        from openpyxl import load_workbook

        from backend.statistics.crew_hours.export import build_crew_hours_workbook

        report, _ = self._member()
        workbook = load_workbook(
            io.BytesIO(
                build_crew_hours_workbook(
                    report, generated_at=datetime(2026, 9, 10, tzinfo=timezone.utc)
                ).getvalue()
            )
        )

        detail = workbook["Cabin"]
        verdicts = {
            detail.cell(row=row, column=6).value: detail.cell(row=row, column=12).value
            for row in (5, 6, 7)
        }
        self.assertEqual(
            verdicts, {"RSX492": "Yes", "RSX121": "No", "RSX122": "No"}
        )
        # The detail block's H.C badge and its block-time total are the values
        # the pre-carve-out export wrote: 1 credit, 10:00 of hours.
        self.assertEqual(detail.cell(row=8, column=12).value, "H.C 1")
        self.assertEqual(detail.cell(row=8, column=11).value, timedelta(hours=10))
        for row in (5, 6, 7):
            self.assertEqual(detail.cell(row=row, column=11).value, time(1, 30), row)

        summary = workbook["Cabin Summary"]
        self.assertEqual(summary.cell(row=2, column=5).value, "H.C")
        self.assertEqual(summary.cell(row=3, column=5).value, 1)   # member H.C
        self.assertEqual(summary.cell(row=4, column=5).value, 1)   # group total
        self.assertEqual(summary.cell(row=3, column=4).value, timedelta(hours=10))
        self.assertEqual(summary.cell(row=4, column=4).value, timedelta(hours=10))


if __name__ == "__main__":
    unittest.main()
